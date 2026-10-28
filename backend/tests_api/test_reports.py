"""Step 7: ปุ่ม "สร้างรายงานสรุปอัตโนมัติด้วย AI" ผ่าน API + Postgres + worker (โหมดทดสอบของ Gemini — คำตอบสังเคราะห์
จาก backend/tests/fixtures/gemini_synthetic.json ไม่ใช่ผล Gemini จริง; container ทดสอบไม่มีอินเทอร์เน็ต)

- การกดหนึ่งครั้ง = job summary → report Word → report PowerPoint ตามลำดับ (resume ทีละงาน) เรียก Gemini ครั้งเดียว
  แล้วหน้าแสดงบทสรุป, ข้อความ "สร้างเมื่อวันที่ …" (เวลาเดียวกับในไฟล์) และปุ่มดาวน์โหลด Word/PowerPoint/.txt ของหน้าเดิม
- ไฟล์ดาวน์โหลดเป็นของ session นั้น; กดซ้ำ = เรียก Gemini ใหม่ (หน้าเดิมไม่เก็บบทสรุปไว้)
- Gemini ล้มเหลว (429) → ข้อความเดิมของหน้า ไม่มีไฟล์; job summary ที่ worker หายไม่ถูกเรียกซ้ำอัตโนมัติ
"""
import io
import json
import re
from pathlib import Path

import pytest
from conftest import ORIGIN, find, nodes, post_event, run_jobs
from docx import Document
from pptx import Presentation
from test_sheets import _research

from app import db, gemini_fixture, summary, worker

FIXTURE = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "gemini_synthetic.json"
SYNTHETIC = json.loads(FIXTURE.read_text(encoding="utf-8"))["responses"][0]["text"]
LABELS = ["📄 ดาวน์โหลดสรุปเป็น Word (มีโลโก้ แก้ไขได้)", "ดาวน์โหลดสไลด์นำเสนอ (.pptx)", "ดาวน์โหลดสรุปเป็น .txt"]


@pytest.fixture()
def gemini(monkeypatch, tmp_path):
    def use(responses=None):
        path = FIXTURE
        if responses is not None:
            path = tmp_path / "gemini.json"
            path.write_text(json.dumps({"note": "synthetic", "responses": responses}), encoding="utf-8")
        monkeypatch.setenv("TFP_GEMINI_FIXTURE", str(path))
        gemini_fixture.CALLS.clear()
    use()
    return use


def _research_home_with_data(client, url):
    csrf = _research(client)
    assert post_event(client, csrf, trigger="nav_top_fetch").json()["pending"]
    run_jobs(url)
    v = post_event(client, csrf, resume=True).json()
    assert find(v, type="button", id="ai_cta_button")
    return csrf


def _jobs(url):
    with db.tx(url) as conn:
        return conn.execute("SELECT kind, status, attempts FROM jobs WHERE kind IN ('summary', 'report') "
                            "ORDER BY created_at").fetchall()


def _pending_kind(url, client):
    """ชนิดของ job ที่หน้ารออยู่ (API บอกแค่ว่ารอ — อ่านจากสถานะ session)"""
    with db.tx(url) as conn:
        return conn.execute("SELECT state->'pending'->>'kind' FROM sessions WHERE token_hash = %s",
                            (db._hash(client.cookies.get("tfp_session")),)).fetchone()[0]


def _click_and_finish(client, csrf, url):
    """กดปุ่ม AI แล้ว resume ทีละงานจนหน้าไม่รอ — คืน (view สุดท้าย, ชนิดงานที่รอตามลำดับ)"""
    v = post_event(client, csrf, trigger="ai_cta_button").json()
    kinds = []
    while v["pending"]:
        kinds.append(_pending_kind(url, client))
        run_jobs(url)
        v = post_event(client, csrf, resume=True).json()
    return v, kinds


def test_ai_summary_runs_once_as_chain_then_shows_original_downloads(client, fresh_db, gemini):
    csrf = _research_home_with_data(client, fresh_db)
    v, kinds = _click_and_finish(client, csrf, fresh_db)
    assert kinds == ["summary", "report", "report"]
    assert len(gemini_fixture.CALLS) == 1                       # resume 3 รอบ เรียก Gemini ครั้งเดียว
    assert [(k, s, a) for k, s, a in _jobs(fresh_db)] == [("summary", "succeeded", 1), ("report", "succeeded", 1),
                                                          ("report", "succeeded", 1)]
    text = json.dumps(v, ensure_ascii=False)
    assert "ข้อความทดสอบสังเคราะห์" in text and "เรียน ท่านผู้บริหาร" not in text  # _web_summary_text ของหน้าเดิม
    assert "| ทดสอบ | 1 |" not in text                                           # ตัดแถวตาราง markdown
    footer = [n for n in nodes(v) if n["type"] == "markdown" and "จัดทำโดยระบบปัญญาประดิษฐ์ Google Gemini" in n["body"]]
    assert len(footer) == 1
    when = re.search(r"สร้างเมื่อวันที่ (\d+ \S+ \d{4} เวลา \d{2}:\d{2} น\.)", footer[0]["body"]).group(1)
    dls = find(v, type="download_button")
    assert [d["label"] for d in dls if d["label"] in LABELS] == LABELS
    files = {d["label"]: client.get(d["url"]) for d in dls if d["label"] in LABELS}
    assert all(r.status_code == 200 for r in files.values())
    word = Document(io.BytesIO(files[LABELS[0]].content))
    assert any(when in p.text for p in word.paragraphs)                   # เวลาในไฟล์ = เวลาที่หน้าแสดง
    slides = Presentation(io.BytesIO(files[LABELS[1]].content)).slides
    assert any(when in sh.text_frame.text for s in slides for sh in s.shapes if sh.has_text_frame)
    txt = files[LABELS[2]].content.decode("utf-8")
    assert txt.endswith(footer[0]["body"].strip()) and "ข้อความทดสอบสังเคราะห์" in txt

    from fastapi.testclient import TestClient

    from app.main import app
    with TestClient(app, base_url=ORIGIN) as other:                       # ไฟล์เป็นของ session นั้น
        other.get("/api/v1/view")
        assert all(other.get(d["url"]).status_code == 404 for d in dls if d["label"] in LABELS)


def test_clicking_again_calls_gemini_again(client, fresh_db, gemini):
    csrf = _research_home_with_data(client, fresh_db)
    _click_and_finish(client, csrf, fresh_db)
    _click_and_finish(client, csrf, fresh_db)
    assert len(gemini_fixture.CALLS) == 2
    assert [k for k, _, _ in _jobs(fresh_db)] == ["summary", "report", "report"] * 2


def test_gemini_quota_error_shows_original_message_and_no_files(client, fresh_db, gemini, monkeypatch):
    e429 = {"error": {"code": 429, "response_json": {"error": {"code": 429, "status": "RESOURCE_EXHAUSTED",
                                                               "message": "synthetic test error (not from Gemini)"}}}}
    gemini([e429])
    waits = []
    monkeypatch.setattr(summary.time, "sleep", waits.append)
    csrf = _research_home_with_data(client, fresh_db)
    v, kinds = _click_and_finish(client, csrf, fresh_db)
    assert kinds == ["summary"] and len(gemini_fixture.CALLS) == 5 and waits == [3.0, 6.0, 12.0, 24.0]
    errors = [n["body"] for n in find(v, type="alert", kind="error")]
    assert any(b.startswith("ขณะนี้มีการเรียกใช้งาน Gemini API ถี่เกินโควตาที่กำหนดไว้ชั่วคราว") for b in errors)
    assert not [d for d in find(v, type="download_button") if d["label"] in LABELS]


def test_lost_summary_job_is_not_retried(client, fresh_db, gemini):
    csrf = _research_home_with_data(client, fresh_db)
    v = post_event(client, csrf, trigger="ai_cta_button").json()
    assert v["pending"] and _pending_kind(fresh_db, client) == "summary"
    job = worker.claim(fresh_db)                         # worker รับงานแล้วหายไป (lease หมด)
    with db.tx(fresh_db) as conn:
        conn.execute("UPDATE jobs SET lease_until = now() - interval '1 second' WHERE id = %s", (job["id"],))
    assert worker.reclaim_expired(fresh_db) == 1
    assert run_jobs(fresh_db) == 0                       # ไม่กลับเข้าคิว
    v = post_event(client, csrf, resume=True).json()
    assert [(k, s) for k, s, _ in _jobs(fresh_db)] == [("summary", "failed")] and not gemini_fixture.CALLS
    assert find(v, type="alert", kind="error", body="เรียก Gemini ไม่สำเร็จ: งานหยุดกลางคัน (worker ไม่ตอบ) กรุณาลองใหม่")
